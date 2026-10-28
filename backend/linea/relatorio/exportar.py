"""Exportadores do relatório: .docx (python-docx), .odt (odfpy) e .pdf (reportlab)."""
from __future__ import annotations

import io
import os
from typing import List

from .conteudo import AVISO, CINZA, Elemento, montar

FORMATOS = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "odt": "application/vnd.oasis.opendocument.text",
    "pdf": "application/pdf",
}
AZUL = "1F3A5F"
_FONTES = os.path.join(os.path.dirname(__file__), "fontes")


def exportar(res: dict, formato: str) -> bytes:
    elementos = montar(res)
    if formato == "docx":
        return _docx(elementos)
    if formato == "odt":
        return _odt(elementos)
    if formato == "pdf":
        return _pdf(elementos)
    raise ValueError(f"Formato não suportado: {formato}")


# =============================================================================
# DOCX
# =============================================================================
def _docx(E: List[Elemento]) -> bytes:
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    util = sec.page_width - sec.left_margin - sec.right_margin
    est = doc.styles["Normal"]
    est.font.name, est.font.size = "Calibri", Pt(10.5)

    def rgb(h):
        return RGBColor.from_string(h) if h else None

    def sombrear(celula, cor_hex):
        tcPr = celula._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), cor_hex)
        tcPr.append(shd)

    for e in E:
        if e.tipo == "titulo":
            p = doc.add_paragraph()
            run = p.add_run(e.texto)
            run.bold, run.font.size, run.font.color.rgb = True, Pt(17), rgb(AZUL)
        elif e.tipo in ("secao", "subsecao"):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12 if e.tipo == "secao" else 9)
            p.paragraph_format.keep_with_next = True
            run = p.add_run(e.texto)
            run.bold = True
            run.font.size = Pt(13 if e.tipo == "secao" else 11.5)
            run.font.color.rgb = rgb(e.cor or AZUL)
        elif e.tipo == "aviso":
            t = doc.add_table(rows=1, cols=1)
            cel = t.cell(0, 0)
            sombrear(cel, "FFF4E5")
            run = cel.paragraphs[0].add_run("Isenção de responsabilidade: ")
            run.bold, run.font.size = True, Pt(9.5)
            run = cel.paragraphs[0].add_run(e.texto)
            run.font.size = Pt(9.5)
            doc.add_paragraph()
        elif e.tipo == "paragrafo":
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            if e.rotulo:
                r0 = p.add_run(e.rotulo)
                r0.bold = True
                if e.cor:
                    r0.font.color.rgb = rgb(e.cor)
            run = p.add_run(e.texto)
            run.bold = e.negrito
            if e.cor:
                run.font.color.rgb = rgb(e.cor)
        elif e.tipo == "lista":
            for it in e.itens:
                p = doc.add_paragraph(style="List Bullet")
                run = p.add_run(it)
                if e.cor:
                    run.font.color.rgb = rgb(e.cor)
        elif e.tipo == "tabela":
            t = doc.add_table(rows=1, cols=len(e.cabecalho))
            t.style = "Table Grid"
            t.alignment = WD_TABLE_ALIGNMENT.CENTER
            for j, h in enumerate(e.cabecalho):
                cel = t.rows[0].cells[j]
                sombrear(cel, "E8EDF3")
                run = cel.paragraphs[0].add_run(h)
                run.bold, run.font.size = True, Pt(9.5)
            for i, linha in enumerate(e.linhas):
                cells = t.add_row().cells
                cor = e.cores_linhas[i] if i < len(e.cores_linhas) else None
                for j, v in enumerate(linha):
                    run = cells[j].paragraphs[0].add_run(v)
                    run.font.size = Pt(9.5)
                    if cor and j > 0:
                        run.font.color.rgb = rgb(cor)
            if e.larguras:
                for row in t.rows:
                    for j, w in enumerate(e.larguras):
                        row.cells[j].width = int(util * w)
            doc.add_paragraph().paragraph_format.space_after = Pt(2)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# =============================================================================
# ODT
# =============================================================================
def _odt(E: List[Elemento]) -> bytes:
    from odf.opendocument import OpenDocumentText
    from odf.style import (ParagraphProperties, Style, TableCellProperties, TableColumnProperties,
                           TableProperties, TextProperties)
    from odf.table import Table, TableCell, TableColumn, TableRow
    from odf.text import List as OList, ListItem, P, Span

    doc = OpenDocumentText()
    estilos = {}

    def est_texto(cor=None, negrito=False, tamanho="10.5pt"):
        chave = ("t", cor, negrito, tamanho)
        if chave not in estilos:
            s = Style(name=f"T{len(estilos)}", family="text")
            kw = {"fontsize": tamanho, "fontfamily": "Calibri"}
            if cor:
                kw["color"] = "#" + cor
            if negrito:
                kw["fontweight"] = "bold"
            s.addElement(TextProperties(**kw))
            doc.automaticstyles.addElement(s)
            estilos[chave] = s
        return estilos[chave]

    def est_par(espaco_antes="0cm", espaco_depois="0.12cm", fundo=None):
        chave = ("p", espaco_antes, espaco_depois, fundo)
        if chave not in estilos:
            s = Style(name=f"P{len(estilos)}", family="paragraph")
            kw = {"margintop": espaco_antes, "marginbottom": espaco_depois}
            if fundo:
                kw.update(backgroundcolor="#" + fundo, padding="0.15cm")
            s.addElement(ParagraphProperties(**kw))
            s.addElement(TextProperties(fontsize="10.5pt", fontfamily="Calibri"))
            doc.automaticstyles.addElement(s)
            estilos[chave] = s
        return estilos[chave]

    cel_cab = Style(name="CelCab", family="table-cell")
    cel_cab.addElement(TableCellProperties(backgroundcolor="#E8EDF3", border="0.5pt solid #9AA5B1",
                                           padding="0.08cm"))
    cel = Style(name="Cel", family="table-cell")
    cel.addElement(TableCellProperties(border="0.5pt solid #9AA5B1", padding="0.08cm"))
    doc.automaticstyles.addElement(cel_cab)
    doc.automaticstyles.addElement(cel)
    n_tab = 0

    def par(texto, cor=None, negrito=False, tamanho="10.5pt", rotulo="", ps=None):
        p = P(stylename=ps or est_par())
        if rotulo:
            p.addElement(Span(stylename=est_texto(cor, True, tamanho), text=rotulo))
        p.addElement(Span(stylename=est_texto(cor, negrito, tamanho), text=texto))
        return p

    for e in E:
        if e.tipo == "titulo":
            doc.text.addElement(par(e.texto, AZUL, True, "17pt"))
        elif e.tipo == "secao":
            doc.text.addElement(par(e.texto, e.cor or AZUL, True, "13pt", ps=est_par("0.4cm", "0.15cm")))
        elif e.tipo == "subsecao":
            doc.text.addElement(par(e.texto, e.cor or AZUL, True, "11.5pt", ps=est_par("0.3cm", "0.1cm")))
        elif e.tipo == "aviso":
            doc.text.addElement(par(e.texto, None, False, "9.5pt", rotulo="Isenção de responsabilidade: ",
                                    ps=est_par("0.1cm", "0.3cm", "FFF4E5")))
        elif e.tipo == "paragrafo":
            doc.text.addElement(par(e.texto, e.cor, e.negrito, rotulo=e.rotulo))
        elif e.tipo == "lista":
            lst = OList()
            for it in e.itens:
                li = ListItem()
                li.addElement(par("• " + it, e.cor))
                lst.addElement(li)
            doc.text.addElement(lst)
        elif e.tipo == "tabela":
            n_tab += 1
            ts = Style(name=f"Tab{n_tab}", family="table")
            ts.addElement(TableProperties(width="17cm", align="center"))
            doc.automaticstyles.addElement(ts)
            tab = Table(name=f"Tabela{n_tab}", stylename=ts)
            larg = e.larguras or [1 / len(e.cabecalho)] * len(e.cabecalho)
            for j, w in enumerate(larg):
                cs = Style(name=f"Tab{n_tab}C{j}", family="table-column")
                cs.addElement(TableColumnProperties(columnwidth=f"{17 * w:.2f}cm"))
                doc.automaticstyles.addElement(cs)
                tab.addElement(TableColumn(stylename=cs))
            tr = TableRow()
            for h in e.cabecalho:
                tc = TableCell(stylename=cel_cab, valuetype="string")
                tc.addElement(par(h, None, True, "9.5pt", ps=est_par("0cm", "0cm")))
                tr.addElement(tc)
            tab.addElement(tr)
            for i, linha in enumerate(e.linhas):
                cor = e.cores_linhas[i] if i < len(e.cores_linhas) else None
                tr = TableRow()
                for j, v in enumerate(linha):
                    tc = TableCell(stylename=cel, valuetype="string")
                    tc.addElement(par(v, cor if j > 0 else None, False, "9.5pt", ps=est_par("0cm", "0cm")))
                    tr.addElement(tc)
                tab.addElement(tr)
            doc.text.addElement(tab)
            doc.text.addElement(P(stylename=est_par()))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# =============================================================================
# PDF
# =============================================================================
def _pdf(E: List[Elemento]) -> bytes:
    from xml.sax.saxutils import escape

    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    from reportlab.platypus import (KeepTogether, ListFlowable, ListItem, Paragraph, SimpleDocTemplate,
                                    Spacer, Table, TableStyle)

    if "DejaVu" not in pdfmetrics.getRegisteredFontNames():
        pdfmetrics.registerFont(TTFont("DejaVu", os.path.join(_FONTES, "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("DejaVu-Bold", os.path.join(_FONTES, "DejaVuSans-Bold.ttf")))
        pdfmetrics.registerFontFamily("DejaVu", normal="DejaVu", bold="DejaVu-Bold",
                                      italic="DejaVu", boldItalic="DejaVu-Bold")

    base = ParagraphStyle("base", fontName="DejaVu", fontSize=9.2, leading=12.5, spaceAfter=3)
    st = {
        "titulo": ParagraphStyle("titulo", parent=base, fontName="DejaVu-Bold", fontSize=16, leading=20,
                                 textColor=colors.HexColor("#" + AZUL), spaceAfter=4),
        "secao": ParagraphStyle("secao", parent=base, fontName="DejaVu-Bold", fontSize=12.5, leading=16,
                                spaceBefore=12, spaceAfter=5, textColor=colors.HexColor("#" + AZUL)),
        "subsecao": ParagraphStyle("subsecao", parent=base, fontName="DejaVu-Bold", fontSize=10.8, leading=14,
                                   spaceBefore=9, spaceAfter=3),
        "celula": ParagraphStyle("celula", parent=base, fontSize=8.4, leading=10.5, spaceAfter=0),
        "aviso": ParagraphStyle("aviso", parent=base, fontSize=8.6, leading=11.5, backColor=colors.HexColor("#FFF4E5"),
                                borderPadding=6, spaceBefore=4, spaceAfter=10),
    }

    def cor(h):
        return colors.HexColor("#" + h) if h else colors.black

    def p(texto, estilo="base", c=None, negrito=False, rotulo=""):
        s = st.get(estilo, base)
        txt = escape(texto)
        if negrito:
            txt = f"<b>{txt}</b>"
        if rotulo:
            txt = f"<b>{escape(rotulo)}</b>{txt}"
        if c:
            txt = f'<font color="#{c}">{txt}</font>'
        return Paragraph(txt, s)

    largura = A4[0] - 4 * cm
    story = []
    for e in E:
        if e.tipo == "titulo":
            story.append(p(e.texto, "titulo"))
        elif e.tipo == "secao":
            story.append(p(e.texto, "secao", e.cor))
        elif e.tipo == "subsecao":
            story.append(p(e.texto, "subsecao", e.cor or AZUL))
        elif e.tipo == "aviso":
            story.append(Paragraph(f"<b>Isenção de responsabilidade: </b>{escape(e.texto)}", st["aviso"]))
        elif e.tipo == "paragrafo":
            story.append(p(e.texto, "base", e.cor, e.negrito, e.rotulo))
        elif e.tipo == "lista":
            story.append(ListFlowable([ListItem(p(it, "base", e.cor), leftIndent=12) for it in e.itens],
                                      bulletType="bullet", start="•", leftIndent=12,
                                      bulletFontName="DejaVu", bulletFontSize=8))
        elif e.tipo == "tabela":
            larg = e.larguras or [1 / len(e.cabecalho)] * len(e.cabecalho)
            dados = [[p(h, "celula", negrito=True) for h in e.cabecalho]]
            for i, linha in enumerate(e.linhas):
                c = e.cores_linhas[i] if i < len(e.cores_linhas) else None
                dados.append([p(v, "celula", c if j > 0 else None) for j, v in enumerate(linha)])
            t = Table(dados, colWidths=[largura * w for w in larg], repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EDF3")),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#9AA5B1")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(t)
            story.append(Spacer(1, 6))

    def rodape(canvas, doc_):
        canvas.saveState()
        canvas.setFont("DejaVu", 7.5)
        canvas.setFillColor(colors.HexColor("#" + CINZA))
        canvas.drawString(2 * cm, 1.2 * cm, "LineaReport · análise automática de apoio ao avaliador")
        canvas.drawRightString(A4[0] - 2 * cm, 1.2 * cm, f"Página {doc_.page}")
        canvas.restoreState()

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=2 * cm,
                            bottomMargin=2 * cm, title="Análise do Relatório do SisDEA", author="LineaReport")
    doc.build(story, onFirstPage=rodape, onLaterPages=rodape)
    return buf.getvalue()
